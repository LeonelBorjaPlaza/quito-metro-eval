"""
vintage_compare.py
Workstream A, step 1 of the move to the 2026-10-04 REMMAQ delivery.
Compares the new delivery with the earlier vintages the pipeline reads, and
writes counts and summary statistics only (no hourly records).

Reading follows the pipeline (01_read_and_merge.R):
  - stations are matched by name, cleaned as standardize_name() does (lower
    case, no spaces, no accents); the new HUM header "Santonio" is mapped to
    sanantonio here (and counted), since the pipeline's cleaner would drop it;
  - every station column is read as numeric whatever the reader would guess:
    numbers stay numbers, text that parses as a number becomes that number,
    other text ("NA") becomes missing; both kinds of text are counted;
  - time stamps are rounded to the nearest hour (round_date(fecha, "hour"));
    rows off the hour and the largest offset are counted; hours duplicated
    after rounding are averaged, as the pipeline does, and counted.
  Negative values are kept here (the pipeline sets negative pollutant values
  to NA after merging); they are counted.

Inputs (read-only):
  new: /home/leonelb/data/quito-metro-eval/air_quality/raw/2026-10-04_remmaq_all_redownload/
  old: data/raw/remmaq/ (the committed symlink, vintages of 2026-04/05)
  gap-fill: .../raw/2026-09-07_secretaria_ambiente_pm25_gapfill/DATOSFALTANTESPM2.5_ANALISISMETRO.xlsx
Outputs: output/local/diagnostics/vintage_2026-10-04/
Run from air_quality/:
  python3 code/local/diagnostics/vintage_compare.py > ../logs/aq_vintage_compare.log 2>&1
"""
import os
import hashlib
import datetime as dt

import numpy as np
import pandas as pd
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))           # air_quality/
STORE = "/home/leonelb/data/quito-metro-eval/air_quality/raw"
NEW = os.path.join(STORE, "2026-10-04_remmaq_all_redownload")
OLD = os.path.join(ROOT, "data", "raw", "remmaq")
GAPFILL = os.path.join(STORE, "2026-09-07_secretaria_ambiente_pm25_gapfill",
                       "DATOSFALTANTESPM2.5_ANALISISMETRO.xlsx")
OUT = os.path.join(ROOT, "output", "local", "diagnostics", "vintage_2026-10-04")
os.makedirs(OUT, exist_ok=True)

USED = ["PM2.5", "CO", "NO2", "SO2", "TMP", "HUM", "VEL", "DIR", "LLU", "RS", "PRE"]  # read by 01
NOT_USED = ["O3", "PM10", "IUV"]
STATIONS = ["belisario", "carapungo", "centro", "cotocollao", "elcamal",
            "guamani", "loschillos", "sanantonio", "tumbaco"]
ALIAS = {"santonio": "sanantonio"}
PEAK = [7, 8, 9, 17, 18, 19]
PERIODS = [("a_before_2022-12", None, "2022-11-30 23:00"),
           ("b_2022-12_to_2023-11", "2022-12-01", "2023-11-30 23:00"),
           ("c_2023-12_to_2025-03", "2023-12-01", "2025-03-31 23:00"),
           ("d_2025-04_on", "2025-04-01", None)]


def clean_name(x):
    if x is None:
        return None
    s = str(x).strip().lower().replace(" ", "")
    s = s.translate(str.maketrans("áéíóúñ", "aeioun"))
    return ALIAS.get(s, s)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def read_book(path, label):
    """Return (hourly DataFrame by cleaned station, inventory dict)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    ws.reset_dimensions()                 # do not trust the stored sheet dimensions
    rows = ws.iter_rows(values_only=True)
    hdr_raw = list(next(rows))
    names = [clean_name(h) for h in hdr_raw[1:]]
    inv = dict(file=os.path.basename(path), vintage=label, sheets=";".join(wb.sheetnames),
               header=" | ".join("" if h is None else str(h) for h in hdr_raw),
               n_columns=len(hdr_raw), unnamed_columns=sum(h is None for h in hdr_raw[1:]),
               aliased=";".join(f"{str(h).strip()}->{clean_name(h)}"
                                for h in hdr_raw[1:] if h is not None
                                and str(h).strip().lower().replace(" ", "") in ALIAS))
    ts, vals = [], []
    n_nondate = n_numtext = n_text = n_neg = n_bool = 0
    max_width, beyond = len(hdr_raw), 0
    offsets = []
    for r in rows:
        f = r[0]
        if isinstance(f, (int, float)) and not isinstance(f, bool):
            f = dt.datetime(1899, 12, 30) + dt.timedelta(days=float(f))
        if not isinstance(f, dt.datetime):
            n_nondate += 1
            continue
        ts.append(f)
        max_width = max(max_width, len(r))
        beyond += sum(1 for v in r[len(hdr_raw):] if v is not None and str(v).strip() != "")
        row = []
        for v in r[1:]:                   # every cell, including any past the header
            if isinstance(v, bool):
                n_bool += 1
                x = np.nan
            elif isinstance(v, (int, float)):
                x = float(v)
            elif isinstance(v, str) and v.strip():
                try:
                    x = float(v.strip())
                    n_numtext += 1
                except ValueError:
                    x = np.nan
                    n_text += 1
            else:
                x = np.nan
            if x == x and x < 0:
                n_neg += 1
            row.append(x)
        vals.append(row)
    wb.close()
    vals = [row + [np.nan] * (max_width - 1 - len(row)) for row in vals]
    t = pd.to_datetime(pd.Series(ts))
    rounded = t.dt.round("h")
    signed = (t - rounded).dt.total_seconds().round(3)
    off = signed.abs()
    cols = [n if n is not None else f"_unnamed{i}" for i, n in enumerate(names)]
    cols += [f"_beyond{j}" for j in range(len(hdr_raw), max_width)]   # cells with no header
    d = pd.DataFrame(vals, columns=cols)
    d.index = rounded.values
    n_dup = int(d.index.duplicated(keep=False).sum())
    d = d.groupby(level=0).mean()          # pipeline averages duplicate date-hours
    inv.update(data_rows=len(t), first=str(rounded.min()), last=str(rounded.max()),
               rows_off_hour=int((off > 0).sum()), rows_off_ge_0_5s=int((off >= 0.5).sum()),
               rows_off_ge_1s=int((off >= 1).sum()), rows_off_gt_1s=int((off > 1).sum()),
               first_off_ge_1s=str(t[off >= 1].min()) if (off >= 1).any() else "",
               max_offset_s=round(float(off.max()), 3),
               rows_ge_1s_after_hour=int((signed >= 1).sum()), rows_ge_1s_before_hour=int((signed <= -1).sum()),
               widest_row_xml_cells=max_width, nonempty_cells_beyond_header=beyond,
               rows_in_duplicated_hours=n_dup, rows_first_cell_not_date=n_nondate,
               numeric_text_cells=n_numtext, nonnumeric_text_cells=n_text,
               negative_values=n_neg, boolean_cells=n_bool,
               stations=";".join(c for c in d.columns if not c.startswith("_")),
               sha256=sha256(path))
    return d, inv


# ------------------------------------------------------------ read everything
books, inventory = {}, []
for var in USED + NOT_USED:
    for label, folder in (("new", NEW), ("old", OLD)):
        p = os.path.join(folder, var + ".xlsx")
        if not os.path.exists(p):
            inventory.append(dict(file=var + ".xlsx", vintage=label, sheets="(not in this vintage)"))
            continue
        d, inv = read_book(p, label)
        books[(var, label)] = d
        inventory.append(inv)
        print(f"read {label} {var}: {inv['data_rows']} rows, {inv['first']} to {inv['last']}, "
              f"off-hour {inv['rows_off_hour']} (max {inv['max_offset_s']} s), dup rows {inv['rows_in_duplicated_hours']}, "
              f"num-text {inv['numeric_text_cells']}, text {inv['nonnumeric_text_cells']}, neg {inv['negative_values']}",
              flush=True)
pd.DataFrame(inventory).to_csv(os.path.join(OUT, "vintage_inventory.csv"), index=False)

# ------------------------------------------------------------ item 2: compare
month_rows, period_rows, year_rows = [], [], []
for var in USED:
    o, n = books[(var, "old")], books[(var, "new")]
    idx = o.index.union(n.index)
    o, n = o.reindex(idx), n.reindex(idx)
    for s in STATIONS:
        if s not in o.columns and s not in n.columns:
            continue
        so = o[s] if s in o.columns else pd.Series(np.nan, index=idx)
        sn = n[s] if s in n.columns else pd.Series(np.nan, index=idx)
        po, pn = so.notna(), sn.notna()
        both = po & pn
        ad = (sn - so).abs()
        differ = both & (ad > 1e-6)
        ratio = (sn / so).where(differ & (so > 0))
        frame = pd.DataFrame(dict(po=po, pn=pn, both=both, differ=differ,
                                  ad=ad.where(differ), ratio=ratio,
                                  new_only=pn & ~po, old_only=po & ~pn), index=idx)

        def summarise(g):
            return pd.Series(dict(
                hours_old=int(g.po.sum()), hours_new=int(g.pn.sum()), hours_both=int(g.both.sum()),
                differ=int(g.differ.sum()),
                median_absdiff=round(float(g.ad.median()), 4) if g.differ.any() else np.nan,
                max_absdiff=round(float(g.ad.max()), 4) if g.differ.any() else np.nan,
                median_ratio_new_old=round(float(g.ratio.median()), 4) if g.ratio.notna().any() else np.nan,
                hours_ratio=int(g.ratio.notna().sum()),                       # differing hours with old > 0
                hours_differ_old_le0=int((g.differ & g.ratio.isna()).sum()),  # differing hours with old <= 0
                new_only=int(g.new_only.sum()), old_only=int(g.old_only.sum())))

        yy = frame.groupby(frame.index.year).apply(summarise)
        yy = yy[(yy.hours_both > 0) & (yy.differ > 0)]
        for y, row in yy.iterrows():
            year_rows.append(dict(variable=var, station=s, year=int(y), hours_both=int(row.hours_both),
                                  differ=int(row.differ), hours_ratio=int(row.hours_ratio),
                                  hours_differ_old_le0=int(row.hours_differ_old_le0),
                                  median_ratio_new_old=row.median_ratio_new_old))
        mm = frame.groupby(frame.index.to_period("M")).apply(summarise)
        mm.insert(0, "station", s)
        mm.insert(0, "variable", var)
        mm.index = mm.index.astype(str)
        month_rows.append(mm.reset_index().rename(columns={"index": "month"}))
        for name, a, b in PERIODS:
            g = frame.loc[(pd.Timestamp(a) if a else frame.index[0]):(pd.Timestamp(b) if b else frame.index[-1])]
            r = summarise(g).to_dict()
            period_rows.append(dict(variable=var, station=s, period=name, **r))
months = pd.concat(month_rows, ignore_index=True)
# Median new/old ratio by station and year (years with differing hours), from the same frame
pd.DataFrame(year_rows).to_csv(os.path.join(OUT, "vintage_ratio_by_station_year.csv"), index=False)
months.to_csv(os.path.join(OUT, "vintage_compare_station_month.csv"), index=False)
periods = pd.DataFrame(period_rows)
periods.to_csv(os.path.join(OUT, "vintage_compare_station_period.csv"), index=False)
print("\n== Revisions to history (changed or newly filled, by period) ==")
print(periods.assign(changed=periods.differ, filled=periods.new_only, dropped=periods.old_only)
      .pivot_table(index=["variable", "station"], columns="period",
                   values=["changed", "filled", "dropped"], aggfunc="sum").to_string())

# ------------------------------------------------------------ item 2e: RS night check
night = [0, 1, 2, 3, 4, 20, 21, 22, 23]
rs_rows = []
for label in ("old", "new"):
    d = books[("RS", label)]
    nd = d[d.index.hour.isin(night)]
    md = d[d.index.hour.isin([10, 11, 12, 13, 14])]
    for s in [c for c in d.columns if c in STATIONS]:
        a = nd[s].groupby(nd.index.to_period("M"))
        b = md[s].groupby(md.index.to_period("M"))
        tab = pd.DataFrame(dict(night_hours=a.count(),
                                night_share_gt10=a.apply(lambda x: (x.dropna() > 10).mean() if len(x.dropna()) else np.nan),
                                night_share_gt50=a.apply(lambda x: (x.dropna() > 50).mean() if len(x.dropna()) else np.nan),
                                night_mean=a.mean().round(1), midday_hours=b.count(), midday_mean=b.mean().round(1)))
        tab.insert(0, "station", s)
        tab.insert(0, "vintage", label)
        tab.index = tab.index.astype(str)
        rs_rows.append(tab.reset_index().rename(columns={"index": "month"}))
rs = pd.concat(rs_rows, ignore_index=True)
rs.to_csv(os.path.join(OUT, "rs_night_check.csv"), index=False)
flag = rs[(rs.night_share_gt10 > 0.2) & (rs.month >= "2024-01")]
print("\n== RS months (2024 on) with more than 20 percent of night hours above 10 W/m2 ==")
print(flag.to_string(index=False) if len(flag) else "none")

# ------------------------------------------------------------ item 3: Los Chillos PM2.5
pn = books[("PM2.5", "new")]["loschillos"]
po = books[("PM2.5", "old")]["loschillos"]
win = pd.date_range("2025-01-06", "2025-03-02 23:00", freq="h")
lc = dict(window="2025-01-06 00:00 to 2025-03-02 23:00", hours_in_window=len(win),
          new_hours=int(pn.reindex(win).notna().sum()), old_hours=int(po.reindex(win).notna().sum()),
          new_hours_days_2025_03_09_to_2025_03_22=int(pn.reindex(pd.date_range("2025-03-09", "2025-03-22 23:00", freq="h")).notna().sum()))
wbg = openpyxl.load_workbook(GAPFILL, read_only=True, data_only=True)
rg = wbg.worksheets[0].iter_rows(values_only=True)
gh = [clean_name(h) for h in next(rg)]
gi = gh.index("loschillos")
g = {}
for r in rg:
    if isinstance(r[0], dt.datetime) and isinstance(r[gi], (int, float)) and not isinstance(r[gi], bool):
        k = pd.Timestamp(r[0]).round("h")
        assert k not in g, f"duplicate hour in gap-fill file: {k}"
        g[k] = float(r[gi])
wbg.close()
g = pd.Series(g).sort_index()
lag_rows = []
for L in range(-3, 4):
    nv = pn.reindex(g.index + pd.Timedelta(hours=L))
    nv.index = g.index
    ok = nv.notna()
    dif = (nv[ok] - g[ok]).abs()
    lag_rows.append(dict(lag_hours_new_file_later_if_positive=L, paired=int(ok.sum()),
                         exact=int((dif < 1e-9).sum()),
                         equal_after_rounding_2dp=int((nv[ok].round(2) == g[ok].round(2)).sum()),
                         within_0_01=int((dif <= 0.01 + 1e-9).sum()),
                         median_absdiff=round(float(dif.median()), 4) if ok.any() else np.nan,
                         max_absdiff=round(float(dif.max()), 4) if ok.any() else np.nan,
                         corr=round(float(np.corrcoef(nv[ok], g[ok])[0, 1]), 4) if ok.sum() > 2 else np.nan))
lags = pd.DataFrame(lag_rows)
lags.to_csv(os.path.join(OUT, "loschillos_new_vs_gapfill.csv"), index=False)
pd.DataFrame([lc]).to_csv(os.path.join(OUT, "loschillos_new_coverage.csv"), index=False)
print("\n== Los Chillos PM2.5 ==", lc)
print(lags.to_string(index=False))

# weekly Los Chillos peak-hour coverage around the March hole (Monday weeks)
pk = pn[(pn.index.dayofweek < 5) & pn.index.hour.isin(PEAK)]
wk = pk.loc["2025-02-17":"2025-04-06"]
wkc = wk.groupby(wk.index.to_period("W-SUN")).count()
wkc.index = [str(p.start_time.date()) for p in wkc.index]
wkc = wkc.reindex([str(d.date()) for d in pd.date_range("2025-02-17", "2025-03-31", freq="7D")], fill_value=0)
wkc.rename("loschillos_weekday_peak_hours").to_csv(os.path.join(OUT, "loschillos_weeks_feb_apr_2025.csv"),
                                                   index_label="week_start")
print(wkc.to_string())

# ------------------------------------------------------------ item 4: timing by month
def morning_peak(s):
    s = s[(s.index.dayofweek < 5) & (s.index.hour >= 4) & (s.index.hour <= 12)]
    grp = s.groupby([s.index.to_period("M"), s.index.hour])
    prof = grp.mean().unstack().reindex(columns=range(4, 13))
    cnt = grp.count().unstack().reindex(columns=range(4, 13)).fillna(0)
    nh = s.groupby(s.index.to_period("M")).count()
    prof = prof[(cnt >= 10).all(axis=1)]                # every hour 04-12 seen on at least 10 weekdays
    pkh = prof.idxmax(axis=1)
    return pkh, nh


def rs_centroid(s):
    s = s.clip(lower=0)
    s = s[(s.index.hour >= 5) & (s.index.hour <= 19)]
    grp = s.groupby([s.index.to_period("M"), s.index.hour])
    g = grp.mean().unstack().reindex(columns=range(5, 20))
    cnt = grp.count().unstack().reindex(columns=range(5, 20)).fillna(0)
    g = g[(cnt >= 10).all(axis=1)]                      # every hour 05-19 seen on at least 10 days
    return (g * g.columns.values).sum(axis=1) / g.sum(axis=1)


t_rows = []
start = "2023-01-01"
for var in ("PM2.5", "CO", "NO2", "SO2", "O3"):
    d = books[(var, "new")].loc[start:]
    for s in [c for c in d.columns if c in STATIONS]:
        pkh, nh = morning_peak(d[s])
        for m in pkh.index:
            t_rows.append(dict(measure=f"{var} weekday morning peak hour (04-12)", station=s,
                               month=str(m), value=pkh[m], hours=int(nh.get(m, 0))))
d = books[("RS", "new")].loc[start:]
for s in [c for c in d.columns if c in STATIONS]:
    c = rs_centroid(d[s])
    for m in c.index:
        t_rows.append(dict(measure="RS daily centroid hour (radiation-weighted mean hour)", station=s,
                           month=str(m), value=round(float(c[m]), 2) if c[m] == c[m] else np.nan, hours=np.nan))
d = books[("TMP", "new")].loc[start:]
for s in [c for c in d.columns if c in STATIONS]:
    gtm = d[s].groupby([d.index.to_period("M"), d.index.hour]).mean().unstack()
    gtm = gtm[gtm.notna().sum(axis=1) >= 20]
    mx = gtm.idxmax(axis=1)
    for m in mx.index:
        t_rows.append(dict(measure="TMP daily maximum hour", station=s, month=str(m), value=mx[m], hours=np.nan))
# Los Chillos against the mean of a fixed set of stations present in every month
# from 2023 (Belisario, Carapungo, Centro, Cotocollao, Tumbaco), PM2.5, by month.
# others.shift(-L, freq="h") gives the others at t+L; a negative best lag means
# Los Chillos's pattern comes later.
REF = ["belisario", "carapungo", "centro", "cotocollao", "tumbaco"]
dfull = books[("PM2.5", "new")]
others = dfull[REF].mean(axis=1)
lag_rows2 = []
d = dfull.loc[start:]
for m, gm in d.groupby(d.index.to_period("M")):
    lcm = gm["loschillos"]
    cors = {}
    for L in range(-4, 5):
        oo = others.shift(-L, freq="h").reindex(gm.index)
        ok = lcm.notna() & oo.notna()
        if ok.sum() > 100:
            cors[L] = float(np.corrcoef(lcm[ok], oo[ok])[0, 1])
    if not cors:
        continue
    best = max(cors, key=cors.get)
    t_rows.append(dict(measure="PM2.5 Los Chillos best lag vs mean of 5 fixed stations (-4..4; negative = Los Chillos later)",
                       station="loschillos", month=str(m), value=best, hours=int(lcm.notna().sum())))
    lag_rows2.append(dict(month=str(m), hours=int(lcm.notna().sum()), best_lag=best,
                          best_lag_at_edge=best in (-4, 4), corr_best=round(cors[best], 4),
                          corr_lag0=round(cors.get(0, np.nan), 4)))
pd.DataFrame(lag_rows2).to_csv(os.path.join(OUT, "loschillos_pm25_lag_by_month.csv"), index=False)

# The same measure for every station, each against the fixed set without itself
lag_rows3 = []
for st in [c for c in STATIONS if c in dfull.columns]:
    ref = [r for r in REF if r != st]
    oth = dfull[ref].mean(axis=1)
    for m, gm in d.groupby(d.index.to_period("M")):
        x = gm[st]
        cors = {}
        for L in range(-4, 5):
            oo = oth.shift(-L, freq="h").reindex(gm.index)
            ok = x.notna() & oo.notna()
            if ok.sum() > 100:
                cors[L] = float(np.corrcoef(x[ok], oo[ok])[0, 1])
        if not cors:
            continue
        best = max(cors, key=cors.get)
        lag_rows3.append(dict(station=st, month=str(m), hours=int(x.notna().sum()), best_lag=best,
                              best_lag_at_edge=best in (-4, 4), corr_best=round(cors[best], 4),
                              corr_lag0=round(cors.get(0, np.nan), 4)))
lag3 = pd.DataFrame(lag_rows3)
lag3.to_csv(os.path.join(OUT, "pm25_lag_by_station_month.csv"), index=False)
print("\n== PM2.5 best lag by station and month (vs fixed stations; negative = station later) ==")
print(lag3.pivot(index="month", columns="station", values="best_lag").to_string())
timing = pd.DataFrame(t_rows)
timing.to_csv(os.path.join(OUT, "timing_by_month.csv"), index=False)
piv = timing[timing.measure.str.startswith("PM2.5 weekday")].pivot(index="month", columns="station", values="value")
print("\n== PM2.5 weekday morning peak hour by month (new vintage) ==")
print(piv.to_string())
piv = timing[timing.measure.str.startswith("RS daily")].pivot(index="month", columns="station", values="value")
print("\n== RS centroid hour by month (new vintage) ==")
print(piv.to_string())

# ------------------------------------------------------------ item 5: coverage by month (new vintage)
cov = []
for var in USED:
    d = books[(var, "new")].loc["2024-10-01":]
    for s in [c for c in d.columns if c in STATIONS]:
        x = d[s]
        if var in ("PM2.5", "CO", "NO2", "SO2"):
            x = x.where(x >= 0)
        pkx = x[(x.index.dayofweek < 5) & x.index.hour.isin(PEAK)]
        a = x.groupby(x.index.to_period("M")).count()
        b = pkx.groupby(pkx.index.to_period("M")).count()
        for m in a.index:
            cov.append(dict(variable=var, station=s, month=str(m), hours=int(a[m]),
                            weekday_peak_hours=int(b.get(m, 0))))
pd.DataFrame(cov).to_csv(os.path.join(OUT, "coverage_new_by_month.csv"), index=False)
print("\nDone. Outputs in", OUT)
